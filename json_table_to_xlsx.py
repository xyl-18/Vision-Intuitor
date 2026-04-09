#!/usr/bin/env python3
"""Convert W&B table JSON file(s) to a single XLSX.

Usage:
    python json_table_to_xlsx.py <input_json_or_dir> [output_xlsx]

Examples:
    python json_table_to_xlsx.py generations_0_026201ef26e5b3cf7fc0.table.json
    python json_table_to_xlsx.py /path/to/val

Behavior:
    - If input is a file: convert that file to one XLSX.
    - If input is a directory: load all .json tables in that directory,
        merge rows, sort by step ascending, and write one XLSX.
"""

from __future__ import annotations

import json
import sys
from collections.abc import Iterable
from pathlib import Path


def fail(msg: str, code: int = 1) -> None:
    print(f"[ERROR] {msg}")
    raise SystemExit(code)


def load_table(path: Path) -> tuple[list[str], list[list[object]]]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(f"Input file not found: {path}")
    except json.JSONDecodeError as e:
        fail(f"Invalid JSON file: {path} ({e})")

    if not isinstance(data, dict):
        fail("JSON root must be an object")

    columns = data.get("columns")
    rows = data.get("data")

    if not isinstance(columns, list) or not all(isinstance(c, str) for c in columns):
        fail("Invalid or missing 'columns' field")
    if not isinstance(rows, list):
        fail("Invalid or missing 'data' field")

    normalized_rows: list[list[object]] = []
    col_count = len(columns)

    for i, row in enumerate(rows, start=1):
        if not isinstance(row, list):
            fail(f"Row {i} is not a list")
        if len(row) < col_count:
            row = row + [""] * (col_count - len(row))
        elif len(row) > col_count:
            row = row[:col_count]
        normalized_rows.append(row)

    return columns, normalized_rows


def discover_json_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        if input_path.suffix.lower() != ".json":
            fail(f"Input file is not a JSON file: {input_path}")
        return [input_path]

    if input_path.is_dir():
        files = sorted(p for p in input_path.iterdir() if p.is_file() and p.suffix.lower() == ".json")
        if not files:
            fail(f"No JSON files found in directory: {input_path}")
        return files

    fail(f"Input path does not exist: {input_path}")
    return []


def rows_to_dicts(columns: list[str], rows: Iterable[list[object]]) -> list[dict[str, object]]:
    out: list[dict[str, object]] = []
    for row in rows:
        out.append({k: v for k, v in zip(columns, row)})
    return out


def parse_step(value: object) -> float:
    if value is None:
        return float("inf")
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return float("inf")
    return float("inf")


def merge_tables(json_paths: list[Path]) -> tuple[list[str], list[list[object]]]:
    all_dict_rows: list[dict[str, object]] = []
    column_order: list[str] = []
    seen: set[str] = set()

    for path in json_paths:
        try:
            columns, rows = load_table(path)
        except SystemExit:
            # Skip non-table JSON files in directory mode, but keep processing.
            print(f"[WARN] Skip non-table JSON: {path.name}")
            continue

        for c in columns:
            if c not in seen:
                seen.add(c)
                column_order.append(c)

        all_dict_rows.extend(rows_to_dicts(columns, rows))

    if not all_dict_rows:
        fail("No valid table rows found in input JSON files")

    # Keep step first for easier reading.
    if "step" in seen:
        final_columns = ["step"] + [c for c in column_order if c != "step"]
    else:
        final_columns = column_order

    # Stable sort by step (small -> large). Rows with missing/non-numeric step go last.
    indexed_rows = list(enumerate(all_dict_rows))
    indexed_rows.sort(key=lambda x: (parse_step(x[1].get("step")), x[0]))

    merged_rows: list[list[object]] = []
    for _, row_dict in indexed_rows:
        merged_rows.append([row_dict.get(col, "") for col in final_columns])

    return final_columns, merged_rows


def to_xlsx(columns: list[str], rows: list[list[object]], output_path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        fail(
            "openpyxl is required. Install it with: pip install openpyxl"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "wandb_table"

    ws.append(columns)
    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"

    # Keep long prompts readable.
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if isinstance(cell.value, str) and len(cell.value) > 120:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Set practical column widths for this table shape.
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        if col_name.startswith("input_") or col_name.startswith("output_"):
            ws.column_dimensions[col_letter].width = 60
        else:
            ws.column_dimensions[col_letter].width = 14

    wb.save(output_path)


def main() -> None:
    if len(sys.argv) < 2:
        fail("Usage: python json_table_to_xlsx.py <input_json_or_dir> [output_xlsx]", code=2)

    input_path = Path(sys.argv[1]).expanduser().resolve()
    if len(sys.argv) >= 3:
        output_path = Path(sys.argv[2]).expanduser().resolve()
    else:
        if input_path.is_dir():
            output_path = (input_path / "merged_tables.xlsx").resolve()
        else:
            output_path = input_path.with_suffix(".xlsx")

    json_paths = discover_json_files(input_path)
    columns, rows = merge_tables(json_paths)
    to_xlsx(columns, rows, output_path)

    print(f"[OK] Wrote: {output_path}")
    print(f"[INFO] JSON files merged: {len(json_paths)}")
    print(f"[INFO] Columns: {len(columns)}, Rows: {len(rows)}")


if __name__ == "__main__":
    main()
